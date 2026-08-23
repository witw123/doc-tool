"""Generate Excel workbooks from files grouped by leaf directories."""

from __future__ import annotations

import io
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_NAME_LENGTH = 31


@dataclass
class ExportResult:
    content: io.BytesIO
    filename: str
    file_count: int
    leaf_count: int
    unmatched_count: int


def _validate_request(path: str, fields: List[str], delimiter: str) -> Tuple[str, List[str]]:
    raw_path = (path or "").strip()
    if not raw_path:
        raise ValueError("文件夹路径不能为空")

    target_path = os.path.normpath(os.path.abspath(raw_path))
    if not os.path.exists(target_path):
        raise FileNotFoundError(f"文件夹路径不存在: {raw_path}")
    if not os.path.isdir(target_path):
        raise NotADirectoryError(f"路径不是文件夹: {raw_path}")

    if not delimiter:
        raise ValueError("文件名分隔符不能为空")

    cleaned_fields = [str(field).strip() for field in fields if str(field).strip()]
    if not cleaned_fields:
        raise ValueError("至少需要配置一个文件名字段")
    if len(set(cleaned_fields)) != len(cleaned_fields):
        raise ValueError("文件名字段不能重复")
    return target_path, cleaned_fields


def _safe_sheet_name(raw_name: str, used_names: set[str]) -> str:
    name = INVALID_SHEET_CHARS.sub("_", raw_name).strip() or "文件夹"
    name = name[:MAX_SHEET_NAME_LENGTH]
    candidate = name
    index = 2
    used_casefolded = {item.casefold() for item in used_names}
    while candidate.casefold() in used_casefolded:
        suffix = f"_{index}"
        candidate = f"{name[:MAX_SHEET_NAME_LENGTH - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate


def _split_filename(filename: str, fields: List[str], delimiter: str) -> Tuple[List[str], str]:
    stem = os.path.splitext(filename)[0]
    parts = stem.split(delimiter)
    matched = len(parts) == len(fields)

    if len(parts) < len(fields):
        values = parts + [""] * (len(fields) - len(parts))
    else:
        values = parts[: len(fields) - 1] + [delimiter.join(parts[len(fields) - 1 :])]
    return values, "匹配" if matched else "不匹配"


def _leaf_rows(target_path: str, fields: List[str], delimiter: str) -> Tuple[Dict[str, List[List[str]]], int, int]:
    grouped: Dict[str, List[List[str]]] = defaultdict(list)
    file_count = 0
    unmatched_count = 0

    for root, dirs, files in os.walk(target_path, topdown=True):
        dirs.sort()
        files.sort()
        if root == target_path or dirs:
            continue

        rel_dir = os.path.relpath(root, target_path).replace("\\", "/")
        folder_name = os.path.basename(root)
        for filename in files:
            full_path = os.path.join(root, filename)
            if not os.path.isfile(full_path):
                continue
            values, status = _split_filename(filename, fields, delimiter)
            grouped[rel_dir].append([rel_dir, folder_name, filename, *values, status])
            file_count += 1
            unmatched_count += status == "不匹配"

    return grouped, file_count, unmatched_count


def _style_sheet(sheet, headers: List[str], rows: List[List[str]]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col_idx, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        width = min(max(max_length + 2, 10), 48)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def build_filename_export(path: str, fields: List[str], delimiter: str, layout: str) -> ExportResult:
    target_path, cleaned_fields = _validate_request(path, fields, delimiter)
    if layout not in {"grouped_sheets", "single_sheet"}:
        raise ValueError("导出结构参数无效")

    grouped, file_count, unmatched_count = _leaf_rows(target_path, cleaned_fields, delimiter)
    if file_count == 0:
        raise ValueError("未找到可导出的叶子文件夹文件")

    headers = ["文件夹相对路径", "文件夹名称", "原始文件名", *cleaned_fields, "匹配状态"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()

    if layout == "single_sheet":
        sheet = workbook.create_sheet(_safe_sheet_name("文件明细", used_names))
        rows = [row for rel_dir in sorted(grouped) for row in grouped[rel_dir]]
        _style_sheet(sheet, headers, rows)
    else:
        for rel_dir in sorted(grouped):
            sheet = workbook.create_sheet(_safe_sheet_name(os.path.basename(rel_dir), used_names))
            _style_sheet(sheet, headers, grouped[rel_dir])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return ExportResult(
        content=output,
        filename="filename_split_export.xlsx",
        file_count=file_count,
        leaf_count=len(grouped),
        unmatched_count=unmatched_count,
    )


def build_filename_preview(path: str, fields: List[str], delimiter: str) -> Tuple[List[str], List[List[str]], int, int, int]:
    """Collect the same leaf-file rows used by Excel export for on-screen preview."""
    target_path, cleaned_fields = _validate_request(path, fields, delimiter)
    grouped, file_count, unmatched_count = _leaf_rows(target_path, cleaned_fields, delimiter)
    if file_count == 0:
        raise ValueError("未找到可预览的叶子文件夹文件")

    headers = ["文件夹相对路径", "文件夹名称", "原始文件名", *cleaned_fields, "匹配状态"]
    rows = [row for rel_dir in sorted(grouped) for row in grouped[rel_dir]]
    return headers, rows, file_count, len(grouped), unmatched_count
