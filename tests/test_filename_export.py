"""Tests for filename-split Excel exports."""

import asyncio
import os
import tempfile
import unittest

from openpyxl import load_workbook
from fastapi import HTTPException

from backend.filename_exporter import build_filename_export, build_filename_preview
from backend.main import export_filename_xlsx, preview_filename
from backend.models import FilenameExportRequest, FilenamePreviewRequest


class TestFilenameExport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = self.temp_dir.name
        os.makedirs(os.path.join(self.root, "group_a", "leaf"))
        os.makedirs(os.path.join(self.root, "group_b", "leaf"))
        with open(os.path.join(self.root, "root_ignored.txt"), "w", encoding="utf-8") as file:
            file.write("ignored")
        for name in ["Project_2026_001.pdf", "Project_2026.pdf", "Project_2026_003_extra.pdf"]:
            with open(os.path.join(self.root, "group_a", "leaf", name), "w", encoding="utf-8") as file:
                file.write("content")
        with open(os.path.join(self.root, "group_b", "leaf", "Other_2025_002.docx"), "w", encoding="utf-8") as file:
            file.write("content")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_grouped_workbook_splits_names_and_ignores_root_files(self):
        result = build_filename_export(self.root, ["项目", "年份", "序号"], "_", "grouped_sheets")
        workbook = load_workbook(result.content)

        self.assertEqual(result.file_count, 4)
        self.assertEqual(result.leaf_count, 2)
        self.assertEqual(result.unmatched_count, 2)
        self.assertEqual(len(workbook.sheetnames), 2)

        rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("文件夹相对路径", "文件夹名称", "原始文件名", "项目", "年份", "序号", "匹配状态"))
        self.assertIn("Project_2026_001.pdf", [row[2] for row in rows[1:]])
        self.assertNotIn("root_ignored.txt", [row[2] for row in rows[1:]])

    def test_single_sheet_merges_extra_segments_into_last_field(self):
        result = build_filename_export(self.root, ["项目", "年份", "编号"], "_", "single_sheet")
        workbook = load_workbook(result.content)
        rows = list(workbook.active.iter_rows(values_only=True))
        extra_row = next(row for row in rows[1:] if row[2] == "Project_2026_003_extra.pdf")
        short_row = next(row for row in rows[1:] if row[2] == "Project_2026.pdf")
        self.assertEqual(extra_row[5], "003_extra")
        self.assertIn(short_row[5], (None, ""))
        self.assertEqual(workbook.sheetnames, ["文件明细"])

    def test_preview_uses_same_leaf_rows_and_dynamic_fields(self):
        headers, rows, file_count, leaf_count, unmatched_count = build_filename_preview(
            self.root,
            ["项目", "年份", "编号"],
            "_",
        )

        self.assertEqual(headers, ["文件夹相对路径", "文件夹名称", "原始文件名", "项目", "年份", "编号", "匹配状态"])
        self.assertEqual(file_count, 4)
        self.assertEqual(leaf_count, 2)
        self.assertEqual(unmatched_count, 2)
        self.assertNotIn("root_ignored.txt", [row[2] for row in rows])
        extra_row = next(row for row in rows if row[2] == "Project_2026_003_extra.pdf")
        self.assertEqual(extra_row[5], "003_extra")

    def test_preview_endpoint_returns_table_data(self):
        response = asyncio.run(preview_filename(FilenamePreviewRequest(
            path=self.root,
            fields=["项目", "年份", "编号"],
            delimiter="_",
        )))

        self.assertTrue(response.success)
        self.assertEqual(response.file_count, 4)
        self.assertEqual(response.leaf_count, 2)
        self.assertEqual(len(response.rows), 4)

    def test_invalid_request_returns_http_error(self):
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(export_filename_xlsx(FilenameExportRequest(
                path=self.root,
                fields=[],
                delimiter="_",
            )))
        self.assertEqual(raised.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()
